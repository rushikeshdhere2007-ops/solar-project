"""
Generate Sign-In Log Excel Sheet
Reads session/login data from solar_storm.db and exports a formatted Excel file.
"""

import os
import sqlite3
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DB_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "solar_storm.db")
OUTPUT_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Sign_In_Log.xlsx")


def fetch_signin_data():
    """Fetch all sign-in records from sessions joined with users."""
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    cursor.execute("""
        SELECT
            s.rowid AS sr_no,
            u.name AS full_name,
            u.username,
            u.role,
            u.level AS clearance_level,
            s.created_at AS signin_datetime
        FROM sessions s
        JOIN users u ON s.user_id = u.id
        ORDER BY s.created_at ASC
    """)
    rows = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return rows


def create_excel(records):
    """Create a professionally formatted Excel sign-in log."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sign-In Log"

    # ── Styles ──────────────────────────────────────────────
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")

    data_font = Font(name="Calibri", size=11)
    alt_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    # ── Title Row ───────────────────────────────────────────
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "🛰️  Solar Storm Mission Control — Sign-In Log"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 38

    # ── Subtitle / Generated timestamp ──────────────────────
    ws.merge_cells("A2:H2")
    subtitle_cell = ws["A2"]
    subtitle_cell.value = f"Generated on: {datetime.now().strftime('%d %B %Y, %I:%M %p')}"
    subtitle_cell.font = Font(name="Calibri", size=10, italic=True, color="555555")
    subtitle_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # ── Headers ─────────────────────────────────────────────
    headers = [
        "Sr. No.",
        "Full Name",
        "Username",
        "Role",
        "Clearance Level",
        "Sign-In Date",
        "Sign-In Time",
        "Day",
    ]
    header_row = 4
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 26

    # ── Data Rows ───────────────────────────────────────────
    for row_idx, record in enumerate(records, start=1):
        excel_row = header_row + row_idx

        # Parse datetime
        try:
            dt = datetime.strptime(record["signin_datetime"], "%Y-%m-%d %H:%M:%S")
        except (ValueError, TypeError):
            dt = None

        row_data = [
            row_idx,
            record.get("full_name", ""),
            record.get("username", ""),
            record.get("role", ""),
            record.get("clearance_level", ""),
            dt.strftime("%d-%b-%Y") if dt else "N/A",
            dt.strftime("%I:%M:%S %p") if dt else "N/A",
            dt.strftime("%A") if dt else "N/A",
        ]

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            cell.font = data_font
            cell.alignment = center if col_idx in (1, 6, 7, 8) else left
            cell.border = thin_border
            # Alternate row shading
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    # ── Column Widths ───────────────────────────────────────
    col_widths = [9, 22, 16, 26, 22, 16, 16, 14]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # ── Summary Footer ──────────────────────────────────────
    footer_row = header_row + len(records) + 2
    ws.merge_cells(f"A{footer_row}:H{footer_row}")
    footer_cell = ws.cell(row=footer_row, column=1)
    footer_cell.value = f"Total Sign-Ins: {len(records)}"
    footer_cell.font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
    footer_cell.alignment = Alignment(horizontal="right", vertical="center")

    # ── Save ────────────────────────────────────────────────
    wb.save(OUTPUT_FILE)
    print(f"[OK] Sign-In Log saved to: {OUTPUT_FILE}")
    print(f"     Total records: {len(records)}")


if __name__ == "__main__":
    records = fetch_signin_data()
    if not records:
        print("[WARNING] No sign-in records found in the database.")
    else:
        create_excel(records)
